import io

import docx
import pytest

from app.parsing import ParsingError, chunk_text, parse_document


def test_parse_txt():
    text = parse_document("lease.txt", b"Hello world, this is a sample lease document with enough text.")
    assert "Hello world" in text


def test_parse_txt_rejects_near_empty_content():
    with pytest.raises(ParsingError):
        parse_document("lease.txt", b"hi")


def test_parse_docx_extracts_paragraphs_and_tables():
    document = docx.Document()
    document.add_paragraph("This is a sample lease clause about rent and deposits.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Rent"
    table.rows[0].cells[1].text = "$2,400"
    buffer = io.BytesIO()
    document.save(buffer)

    text = parse_document("lease.docx", buffer.getvalue())
    assert "sample lease clause" in text
    assert "Rent" in text and "2,400" in text


def test_parse_docx_rejects_corrupt_bytes():
    with pytest.raises(ParsingError):
        parse_document("lease.docx", b"not a real docx file, just garbage bytes" * 5)


def test_parse_unsupported_extension_raises():
    with pytest.raises(ParsingError):
        parse_document("lease.rtf", b"some content that is long enough to pass the length check")


def test_chunk_text_respects_max_chars_and_preserves_content():
    text = "\n".join(f"Paragraph number {i} with some filler words to pad it out a bit." for i in range(40))
    chunks = chunk_text(text, max_chars=200, overlap=20)
    assert len(chunks) > 1
    assert all(len(c) <= 220 for c in chunks)  # small slack for hard-split overlap
    assert "Paragraph number 0" in chunks[0]


def test_chunk_text_single_short_paragraph():
    chunks = chunk_text("Just one short paragraph.", max_chars=800)
    assert chunks == ["Just one short paragraph."]
